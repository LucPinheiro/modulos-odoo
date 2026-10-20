#!/usr/bin/env python3
"""Traduce un Excel exportado de Odoo para importarlo con el módulo.
Uso: GOOGLE_TRANSLATE_API_KEY=... python traducir_excel_google.py --input productos.xlsx --output traducciones.xlsx --source de --target es_ES
Dependencias: pip install openpyxl requests
La API key no se escribe ni se guarda en el archivo generado.
"""
import argparse
import html
import os
import time
from pathlib import Path
import requests
from openpyxl import load_workbook, Workbook

FIELDS = ("short_description_product", "full_description_product")

def translate(text, source, target, api_key, session):
    # Google Cloud Translation API v2 usa el código base es, en, pt, etc.
    target_api = target.split("_")[0].lower()
    source_api = source.split("_")[0].lower() if source else None
    payload = {"q": text, "target": target_api, "format": "html"}
    if source_api: payload["source"] = source_api
    response = session.post(
        "https://translation.googleapis.com/language/translate/v2",
        params={"key": api_key}, json=payload, timeout=60)
    response.raise_for_status()
    return html.unescape(response.json()["data"]["translations"][0]["translatedText"])

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--source", default="", help="Idioma original, p.ej. de_DE; vacío autodetecta")
    ap.add_argument("--target", required=True, help="Idioma de Odoo activo, p.ej. es_ES")
    ap.add_argument("--delay", type=float, default=0.05)
    args = ap.parse_args()
    key = os.getenv("GOOGLE_TRANSLATE_API_KEY")
    if not key: ap.error("Configura GOOGLE_TRANSLATE_API_KEY en una variable de entorno")
    if Path(args.input).resolve() == Path(args.output).resolve(): ap.error("El Excel de salida debe ser diferente del origen")
    wb = load_workbook(args.input, read_only=True, data_only=True)
    sheet = wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
    if "id" not in header or not set(FIELDS).intersection(header):
        ap.error("El Excel necesita columnas id y al menos una descripción")
    if len(header) != len(set(header)): ap.error("Hay cabeceras duplicadas")
    out = Workbook(); ws = out.active; ws.title = "Traducciones"
    ws.append(["id", "lang", *FIELDS])
    session = requests.Session(); count = 0
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(header, values))
        if not row.get("id"): continue
        result = []
        for field in FIELDS:
            source_text = row.get(field)
            if not isinstance(source_text, str) or not source_text.strip():
                result.append(None); continue
            try:
                result.append(translate(source_text, args.source, args.target, key, session))
            except (requests.RequestException, KeyError, IndexError, ValueError) as exc:
                raise SystemExit(f"Error en fila {number}, {field}: {exc}") from exc
            time.sleep(args.delay)
        if any(v is not None for v in result):
            ws.append([row["id"], args.target, *result]); count += 1
            print(f"Procesada fila {number} | ID {row['id']}")
    out.save(args.output)
    wb.close()
    print(f"Generado {args.output} con {count} productos para {args.target}")

if __name__ == "__main__":
    main()
